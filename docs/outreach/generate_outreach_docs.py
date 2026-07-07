#!/usr/bin/env python3
"""Generate DanyPathMart outreach Word documents (general B2B, not school-specific)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = Path(__file__).resolve().parent
BRAND_GREEN = RGBColor(0x1B, 0x5E, 0x20)
BRAND_GOLD = RGBColor(0xC9, 0xA2, 0x27)
MUTED = RGBColor(0x55, 0x55, 0x55)
PLACEHOLDER = RGBColor(0xB8, 0x6E, 0x00)  # amber — stands out in Word


def ph(text: str) -> str:
    """Standard placeholder token for your details (fill once)."""
    return f"[{text}]"


def recipient(text: str) -> str:
    """Per-recipient placeholder (change every send)."""
    return f"[{text}]"


def insert_visual(text: str) -> str:
    """Visual / file insert marker (logo, screenshot, etc.)."""
    return f"◆ INSERT: {text} ◆"


def add_placeholder_legend(doc: Document) -> None:
    heading(doc, "How to read placeholders in this pack", level=2)
    p = doc.add_paragraph()
    r = p.add_run(insert_visual("EXAMPLE"))
    r.bold = True
    r.font.color.rgb = PLACEHOLDER
    p.add_run(" — Drop in a logo, photo, or screenshot in Word (Insert → Pictures). "
              "Delete the ◆ marker text after adding the image.")
    p = doc.add_paragraph()
    r = p.add_run(ph("YOUR FULL NAME"))
    r.bold = True
    r.font.color.rgb = PLACEHOLDER
    p.add_run(" — Your company details. Fill once, reuse everywhere.")
    p = doc.add_paragraph()
    r = p.add_run(recipient("Recipient name"))
    r.bold = True
    r.font.color.rgb = PLACEHOLDER
    p.add_run(" — Change for each person or business you contact.")
    doc.add_paragraph()


def add_insert_box(doc: Document, label: str, hint: str = "") -> None:
    """Visible box showing where to put logos, screenshots, etc."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "FFF8E1")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(insert_visual(label))
    run.bold = True
    run.font.color.rgb = PLACEHOLDER
    run.font.size = Pt(11)
    if hint:
        p2 = cell.add_paragraph(hint)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p2.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = MUTED
            run.italic = True
    doc.add_paragraph()


def add_word_layout_notes(doc: Document) -> None:
    heading(doc, "Word layout — where to put logos & headers", level=2)
    body(doc, "Do these once in Word after opening any document:", bold=True)
    numbered(doc, "Cover logo: replace the yellow ◆ INSERT: LOGO ◆ box (Insert → Pictures → select your PNG/SVG).")
    numbered(doc, "Every page header: double-click top margin → Insert → Pictures → small logo left; type document title right.")
    numbered(doc, "Every page footer: Insert → Page Number (bottom centre) + type " + ph("YOUR WEBSITE URL") + " on the right.")
    numbered(doc, "Brand colours (optional): select headings → Font colour → your green/gold.")
    numbered(doc, "One-pager / partnership pack: add " + insert_visual("DASHBOARD SCREENSHOT") + " and "
             + insert_visual("SAMPLE PRODUCT LISTING") + " in the appendix section.")
    numbered(doc, "Before sending: File → Save As → PDF for WhatsApp/email attachments.")
    doc.add_paragraph()

def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_cover(doc: Document, title: str, subtitle: str, version: str = "v1 — June 2025") -> None:
    add_insert_box(
        doc,
        "LOGO HERE (company logo, PNG recommended, ~300px wide)",
        "Word: click here → Insert → Pictures → This Device",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(ph("COMPANY / BRAND NAME"))
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = PLACEHOLDER

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(18)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run(subtitle)
    sr.font.size = Pt(12)
    sr.font.color.rgb = MUTED

    v = doc.add_paragraph()
    v.alignment = WD_ALIGN_PARAGRAPH.CENTER
    vr = v.add_run(version + "  ·  " + ph("DOCUMENT DATE"))
    vr.font.size = Pt(10)
    vr.font.color.rgb = MUTED

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = contact.add_run(
        ph("YOUR FULL NAME") + "  ·  " + ph("YOUR PHONE") + "  ·  " + ph("YOUR EMAIL")
    )
    cr.font.size = Pt(10)
    cr.font.color.rgb = PLACEHOLDER

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Replace amber [BRACKET] text with your details. Replace ◆ INSERT ◆ boxes with logos or images. "
        "See 00_DPM_Where_To_Put_Your_Details.docx for the full list."
    )
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(10)
    note.runs[0].font.color.rgb = MUTED
    doc.add_page_break()


def heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        if level == 1:
            run.font.color.rgb = BRAND_GREEN


def body(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph(text)
    if bold and p.runs:
        p.runs[0].bold = True


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def numbered(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "E8F5E9")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()


def email_block(doc: Document, title: str, subject: str, body_lines: list[str]) -> None:
    heading(doc, title, level=2)
    p = doc.add_paragraph()
    p.add_run("Subject: ").bold = True
    p.add_run(subject)
    for line in body_lines:
        doc.add_paragraph(line)
    doc.add_paragraph()


def build_customization_guide_doc() -> Path:
    doc = Document()
    add_cover(
        doc,
        "Where To Put Your Details",
        "Master checklist — logos, links, fees, and per-send fields",
    )
    add_placeholder_legend(doc)

    heading(doc, "1. Visual inserts (◆ INSERT ◆)")
    add_table(
        doc,
        ["Location", "What to insert", "How in Word"],
        [
            ["Cover of every document", insert_visual("LOGO HERE"), "Insert → Pictures in the yellow box"],
            ["Header (all pages)", insert_visual("SMALL LOGO + DOC TITLE"), "Double-click header area → Insert → Pictures"],
            ["Partnership pack — cover", insert_visual("LOGO HERE"), "Same as cover box"],
            ["Partnership pack — appendix", insert_visual("DASHBOARD SCREENSHOT"), "Screenshot of seller dashboard"],
            ["Partnership pack — appendix", insert_visual("SAMPLE PRODUCT LISTING"), "Screenshot of a live or mock listing"],
            ["One-pager — top", insert_visual("LOGO HERE"), "Small logo above headline"],
            ["Social post (optional)", insert_visual("PROMO IMAGE OR BANNER"), "Export from Canva/Figma if used"],
        ],
    )

    heading(doc, "2. Your details — fill once, reuse everywhere")
    add_table(
        doc,
        ["Placeholder", "Example", "Used in"],
        [
            [ph("COMPANY / BRAND NAME"), "DanyPathMart", "Covers, emails, one-pager"],
            [ph("YOUR FULL NAME"), "Jane Doe", "Signature, contact lines"],
            [ph("YOUR JOB TITLE"), "Partnerships Lead", "Email signature"],
            [ph("YOUR PHONE"), "+233 XX XXX XXXX", "Signature, WhatsApp, one-pager"],
            [ph("YOUR EMAIL"), "you@company.com", "Signature"],
            [ph("YOUR WHATSAPP NUMBER"), "same or different", "WhatsApp templates"],
            [ph("YOUR WEBSITE URL"), "https://danypathmart.store", "Footer, emails"],
            [ph("SHOP APPLY URL"), "https://danypathmart.store/shop/apply", "All CTAs"],
            [ph("APPLY URL WITH REF CODE"), "…/shop/apply?ref=CODE123", "When sending a referral link"],
            [ph("DOCUMENT DATE"), "June 2025", "Cover, partnership pack"],
        ],
    )

    heading(doc, "3. Offers & fees — fill from your admin settings")
    add_table(
        doc,
        ["Placeholder", "Example", "Where set"],
        [
            [ph("STANDARD REGISTRATION FEE"), "GHS 300", "Admin → Marketplace → Billing"],
            [ph("FREE REGISTRATION UNTIL DATE"), "31 July 2025", "Registration promos section"],
            [ph("FIRST REG DISCOUNTED FEE"), "GHS 250", "First-registration discount"],
            [ph("REFERRAL APPLICANT DISCOUNT"), "10% off or GHS 30 off", "Referral discount section"],
            [ph("REFERRAL / PROMOTER CODE"), "SHOP-ABC or promoter code", "Per partner you assign"],
            [ph("PROMOTER COMMISSION DESCRIPTION"), "X% of registration fee", "Promoter brief & email 3"],
            [ph("RENEWAL FEE AND PERIOD"), "GHS X / month or year", "Partnership pack fees section"],
            [ph("APPROVAL TIMEFRAME"), "3–5 business days", "FAQ answers"],
            [ph("PRODUCT CATEGORIES YOU ACCEPT"), "Uniforms, books, …", "FAQ & one-pager"],
        ],
    )

    heading(doc, "4. Per recipient — change every send")
    add_table(
        doc,
        ["Placeholder", "When to change"],
        [
            [recipient("Recipient name"), "Every email / WhatsApp"],
            [recipient("Business name"), "Every outreach"],
            [recipient("City or area"), "When relevant"],
            [recipient("Product category they sell"), "Personalise cold open"],
            [recipient("Mutual contact name"), "Introduction emails only"],
            [recipient("Key benefit discussed"), "Meeting leave-behind only"],
            [recipient("Follow-up date"), "Leave-behind & reminders"],
        ],
    )

    add_word_layout_notes(doc)

    heading(doc, "5. Quick workflow")
    numbered(doc, "Open 00 (this file) and fill a copy of the tables in a notes doc or on paper.")
    numbered(doc, "Open 02 Email Templates → replace signature block and all " + ph("YOUR…") + " tokens.")
    numbered(doc, "Open 03 Business Samples → build partnership pack; add logos and screenshots.")
    numbered(doc, "Export PDF from 03 one-pager for WhatsApp attachments.")
    numbered(doc, "Use 04 for short messages; 05 to track sends.")

    path = OUT_DIR / "00_DPM_Where_To_Put_Your_Details.docx"
    doc.save(path)
    return path


def build_strategy_doc() -> Path:
    doc = Document()
    add_cover(
        doc,
        "Outreach Strategy & Playbook",
        "How to get sellers and partners on board — general B2B guide",
    )

    heading(doc, "1. Purpose of this pack")
    add_placeholder_legend(doc)
    body(
        doc,
        "This playbook helps you recruit businesses onto the "
        + ph("COMPANY / BRAND NAME")
        + " marketplace. It is intentionally general: you choose the recipient and personalise the opening line.",
    )

    heading(doc, "2. Who you can target")
    bullet(doc, "Retail shops and suppliers (uniforms, books, stationery, club merchandise, general goods)")
    bullet(doc, "Wholesale / distribution businesses looking for an online sales channel")
    bullet(doc, "Individual entrepreneurs and side hustlers with inventory to sell")
    bullet(doc, "Promoters and referral partners who introduce new sellers")
    bullet(doc, "Any organisation or contact in your network who can recommend vendors")

    heading(doc, "3. Three outreach lanes")
    heading(doc, "Lane A — Warm network (fastest)", level=2)
    body(doc, "Start with people you already know or who know you.")
    numbered(doc, "List 20 contacts: former colleagues, suppliers, business owners in your area.")
    numbered(doc, "Send a short personal message (email or WhatsApp) — not a mass blast.")
    numbered(doc, "Offer a 15-minute onboarding call or a direct apply link with a referral code.")
    numbered(doc, "Ask: “Who else should I speak to?” after each yes.")

    heading(doc, "Lane B — Referral & promoter loop (scalable)", level=2)
    body(
        doc,
        "Use DanyPathMart shop codes and promoter codes. Early partners earn when they refer new sellers. "
        "Combine with admin promos: free registration period, first-registration discount, referral applicant discount.",
    )
    add_table(
        doc,
        ["Tool", "What it does", "Your action"],
        [
            ["Free registration until " + ph("FREE REGISTRATION UNTIL DATE"), "Removes signup fee barrier", "Admin → Marketplace → Billing → Registration promos"],
            ["First-registration discount", "Lower fee for brand-new sellers", "e.g. " + ph("STANDARD REGISTRATION FEE") + " → " + ph("FIRST REG DISCOUNTED FEE")],
            ["Referral applicant discount", "Discount when apply with valid code", ph("REFERRAL APPLICANT DISCOUNT")],
            ["Promoter commission", "Referrer earns on paid registration", ph("PROMOTER COMMISSION DESCRIPTION")],
        ],
    )

    heading(doc, "Lane C — Outbound (cold / semi-cold)", level=2)
    body(doc, "Identify businesses that fit your marketplace categories. Research name, city, and one detail you can mention.")
    numbered(doc, "Send Email 1 (see Email Templates document).")
    numbered(doc, "Follow up on WhatsApp or phone after 3–5 days if no reply.")
    numbered(doc, "Send Email 2 (follow-up) once.")
    numbered(doc, "Move on after two touches unless they engage.")

    heading(doc, "4. What converts sellers")
    add_table(
        doc,
        ["Message", "Why it works"],
        [
            ["Low or zero signup risk", "Free period or discounted first registration"],
            ["Clear job description", "“You list products; we bring buyers; you fulfill orders”"],
            ["Fast payout story", "Parents/customers pay online; seller gets MoMo/bank payout"],
            ["Referral incentive", "“Your code gives them a discount; you earn when they join”"],
            ["Human onboarding", "15-minute walkthrough beats a long form alone"],
        ],
    )

    heading(doc, "5. 30-day rollout plan")
    add_table(
        doc,
        ["Week", "Focus", "Target"],
        [
            ["Week 1", "Pilot partners (5–10)", "Warm network only"],
            ["Week 2", "Document early wins", "1–2 short testimonials or order counts"],
            ["Week 3", "Activate referrers", "Give each pilot a unique code"],
            ["Week 4", "Outbound at scale", "20 emails + 20 WhatsApp messages"],
        ],
    )

    heading(doc, "6. Personalisation checklist (before every send)")
    bullet(doc, "Recipient name and business name correct")
    bullet(doc, "One specific line: why them (location, product type, mutual contact)")
    bullet(doc, "Offer matches what admin has enabled (free period, discount, code)")
    bullet(doc, "Single clear call-to-action (one link or one meeting request)")
    bullet(doc, "Your name, role, phone, and " + ph("SHOP APPLY URL") + " in signature")

    heading(doc, "7. Metrics to track")
    bullet(doc, "Messages sent / opened / replied")
    bullet(doc, "Apply link clicks")
    bullet(doc, "Applications started vs completed vs paid")
    bullet(doc, "Shops approved and first product listed")
    bullet(doc, "Referral code usage")

    heading(doc, "8. Folder structure (recommended)")
    body(doc, "Keep files organised under docs/outreach/:")
    bullet(doc, "00 — Where to put your details (start here)")
    bullet(doc, "01 — Strategy (this document)")
    bullet(doc, "02 — Email templates")
    bullet(doc, "03 — Business document samples (partnership pack, one-pager)")
    bullet(doc, "04 — WhatsApp & social short messages")
    bullet(doc, "PDF exports for sharing; DOCX for editing")

    path = OUT_DIR / "01_DPM_Outreach_Strategy_and_Playbook.docx"
    doc.save(path)
    return path


def build_email_templates_doc() -> Path:
    doc = Document()
    add_cover(
        doc,
        "Email Templates",
        "General B2B outreach — personalise [brackets] for each recipient",
    )

    heading(doc, "How to use these templates")
    add_placeholder_legend(doc)
    body(
        doc,
        "Every template is general on purpose. Add one personalised opening sentence naming "
        + recipient("Recipient name")
        + " and "
        + recipient("Business name")
        + ". Do not send unchanged bulk mail.",
    )
    body(doc, "Signature block — copy this to the bottom of every email:", bold=True)
    sig_lines = [
        ph("YOUR FULL NAME"),
        ph("YOUR JOB TITLE") + ", " + ph("COMPANY / BRAND NAME"),
        ph("YOUR PHONE") + " | " + ph("YOUR EMAIL"),
        ph("YOUR WEBSITE URL"),
        "Apply: " + ph("SHOP APPLY URL") + "  (add ?ref=" + ph("REFERRAL / PROMOTER CODE") + " when using a code)",
    ]
    for line in sig_lines:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.color.rgb = PLACEHOLDER

    email_block(
        doc,
        "Template 1 — Cold introduction (any business)",
        "Grow your sales online with " + ph("COMPANY / BRAND NAME"),
        [
            "Dear " + recipient("Recipient name") + ",",
            "",
            "I’m " + ph("YOUR FULL NAME") + " from " + ph("COMPANY / BRAND NAME")
            + " — a Ghana marketplace where businesses list products and reach customers who want to shop and pay online.",
            "",
            "I came across " + recipient("Business name") + " / your work in "
            + recipient("Product category")
            + " and thought you might be a good fit for our seller programme. You keep your brand and pricing; we provide the storefront, checkout, and order flow.",
            "",
            "Right now we’re offering [choose one: free registration until "
            + ph("FREE REGISTRATION UNTIL DATE")
            + " / first registration at "
            + ph("FIRST REG DISCOUNTED FEE")
            + " / referral discount with code "
            + ph("REFERRAL / PROMOTER CODE")
            + "]. Setup takes about 15 minutes.",
            "",
            "Would you be open to a short call this week, or would you prefer to apply directly here?",
            ph("SHOP APPLY URL"),
            "",
            "Best regards,",
            "[paste signature block above]",
        ],
    )

    email_block(
        doc,
        "Template 2 — Warm introduction (someone you know)",
        "Quick idea for " + recipient("Business name") + " — sell on " + ph("COMPANY / BRAND NAME"),
        [
            "Hi " + recipient("Recipient name") + ",",
            "",
            "Hope you’re well. I’m reaching out because we’ve opened "
            + ph("COMPANY / BRAND NAME")
            + " for more sellers, and I immediately thought of "
            + recipient("Business name") + ".",
            "",
            "It’s straightforward: list your products, receive orders, fulfill them, and get paid to your MoMo or bank. We’re running ["
            + ph("CURRENT PROMO SUMMARY")
            + "] for new shops this month.",
            "",
            "If you’re interested, I can walk you through it in 15 minutes or send you "
            + ph("APPLY URL WITH REF CODE") + ".",
            "",
            "Let me know what works for you.",
            "",
            "[paste signature block above]",
        ],
    )

    email_block(
        doc,
        "Template 3 — Referral / promoter invitation",
        "Earn when you refer sellers to " + ph("COMPANY / BRAND NAME"),
        [
            "Dear " + recipient("Recipient name") + ",",
            "",
            ph("COMPANY / BRAND NAME")
            + " is expanding our seller network, and we’re inviting trusted partners to refer businesses onto the platform.",
            "",
            "When someone registers using your code "
            + ph("REFERRAL / PROMOTER CODE")
            + ", they receive "
            + ph("REFERRAL APPLICANT DISCOUNT")
            + ", and you earn "
            + ph("PROMOTER COMMISSION DESCRIPTION")
            + " on their paid registration.",
            "",
            "This suits anyone with a network of shop owners, suppliers, or entrepreneurs — no need to run a shop yourself unless you want to.",
            "",
            "If you’d like to join as a promoter, reply to this email and I’ll set up your account and code.",
            "",
            "[paste signature block above]",
        ],
    )

    email_block(
        doc,
        "Template 4 — Follow-up (no reply after 3–5 days)",
        "Re: " + ph("COMPANY / BRAND NAME") + " seller opportunity",
        [
            "Hi " + recipient("Recipient name") + ",",
            "",
            "Just following up on my note about listing on " + ph("COMPANY / BRAND NAME") + ".",
            "",
            "Brief recap: you can sell online through our marketplace, receive orders, and get paid via MoMo or bank. New sellers currently get "
            + ph("CURRENT PROMO SUMMARY") + ".",
            "",
            "Happy to answer questions by reply or WhatsApp (" + ph("YOUR WHATSAPP NUMBER") + "). Apply: "
            + ph("SHOP APPLY URL"),
            "",
            "Thanks,",
            "[paste signature block above]",
        ],
    )

    email_block(
        doc,
        "Template 5 — After they show interest",
        "Next steps — your " + ph("COMPANY / BRAND NAME") + " shop application",
        [
            "Hi " + recipient("Recipient name") + ",",
            "",
            "Great speaking with you. Here’s what to do next:",
            "",
            "1. Apply here: " + ph("APPLY URL WITH REF CODE"),
            "2. Have ready: business name, contact details, city, and payout info (bank or MoMo)",
            "3. After approval, add your first products with photos and prices",
            "",
            "Registration fee: "
            + ph("STANDARD REGISTRATION FEE")
            + " (or FREE until "
            + ph("FREE REGISTRATION UNTIL DATE")
            + "). With code "
            + ph("REFERRAL / PROMOTER CODE")
            + ", fee is [calculated discounted amount].",
            "",
            "Reply if you hit any issues — I’m happy to help.",
            "",
            "[paste signature block above]",
        ],
    )

    email_block(
        doc,
        "Template 6 — Introduction (you forward to someone they know)",
        "Introduction — " + ph("COMPANY / BRAND NAME") + " marketplace partnership",
        [
            "Dear " + recipient("Recipient name") + ",",
            "",
            recipient("Mutual contact name")
            + " suggested I reach out regarding an opportunity to sell through "
            + ph("COMPANY / BRAND NAME") + ".",
            "",
            ph("COMPANY / BRAND NAME")
            + " is a Ghana-based marketplace. We help businesses reach customers online with integrated checkout and order management. Sellers list products, fulfill orders, and receive payouts to MoMo or bank.",
            "",
            "We’re currently onboarding new sellers with " + ph("CURRENT PROMO SUMMARY") + ". I’d welcome a conversation at your convenience.",
            "",
            "Apply: " + ph("SHOP APPLY URL") + " | Questions: " + ph("YOUR PHONE"),
            "",
            "[paste signature block above]",
        ],
    )

    heading(doc, "Subject line bank (mix and match)")
    bullets = [
        "Sell online with " + ph("COMPANY / BRAND NAME") + " — " + ph("CURRENT PROMO SUMMARY"),
        "Partnership opportunity — " + ph("COMPANY / BRAND NAME") + " marketplace",
        "Quick question about " + recipient("Business name") + " and online sales",
        "Invitation to list on " + ph("COMPANY / BRAND NAME"),
        recipient("Recipient name") + ", thought of you for our seller programme",
    ]
    for b in bullets:
        bullet(doc, b)

    heading(doc, "Email best practices")
    bullet(doc, "One call-to-action per email")
    bullet(doc, "Under 150 words for cold outreach; longer only after they reply")
    bullet(doc, "No attachment on first cold email — link to apply or attach PDF on follow-up")
    bullet(doc, "Send Tue–Thu, 9–11am for business email; WhatsApp often gets faster replies in Ghana")

    path = OUT_DIR / "02_DPM_Email_Templates.docx"
    doc.save(path)
    return path


def build_business_samples_doc() -> Path:
    doc = Document()
    add_cover(
        doc,
        "Business Document Samples",
        "Partnership pack, one-pager, and leave-behind materials for prospects",
    )

    heading(doc, "Document A — Seller Partnership Pack (4–6 pages)")
    body(doc, "Build this as a separate Word file when sending a formal PDF. Use the markers below.")
    add_insert_box(doc, "LOGO HERE — partnership pack cover", "Top of page 1, centred")
    sections = [
        ("Cover", insert_visual("LOGO HERE") + " | Seller Partnership Pack | " + ph("DOCUMENT DATE") + " | " + ph("YOUR PHONE") + " | " + ph("SHOP APPLY URL")),
        ("1. Executive summary (½ page)", ph("COMPANY / BRAND NAME") + " — what it is. Who can sell. Why join now (" + ph("CURRENT PROMO SUMMARY") + ")."),
        (
            "2. How it works",
            "Apply → Review → List products → Receive orders → Fulfill → Get paid. Number each step.",
        ),
        (
            "3. Roles & responsibilities (table)",
            "Seller vs " + ph("COMPANY / BRAND NAME") + " — see sample table below.",
        ),
        (
            "4. Fees & current promotions",
            ph("STANDARD REGISTRATION FEE") + ", " + ph("RENEWAL FEE AND PERIOD") + ", "
            + ph("FREE REGISTRATION UNTIL DATE") + ", " + ph("REFERRAL APPLICANT DISCOUNT"),
        ),
        ("5. Requirements", "Business name, email, phone, city, product photos, payout details."),
        ("6. FAQ", "Use sample answers at end of this document."),
        ("7. Call to action", ph("SHOP APPLY URL") + " | " + ph("YOUR WHATSAPP NUMBER") + " | code field on apply form"),
        ("Appendix", insert_visual("DASHBOARD SCREENSHOT") + " | " + insert_visual("SAMPLE PRODUCT LISTING") + " | " + ph("YOUR WEBSITE URL") + "/terms"),
    ]
    add_table(doc, ["Section", "Content — replace placeholders"], sections)

    heading(doc, "Document B — One-pager (1 page)")
    body(doc, "Copy into a new Word file. Keep to one page when exported as PDF.")
    add_insert_box(doc, "LOGO HERE — one-pager (small, top centre)", "~150px wide recommended")
    doc.add_paragraph()
    one_pager = doc.add_table(rows=9, cols=1)
    one_pager.style = "Table Grid"
    lines = [
        ph("COMPANY / BRAND NAME").upper() + " — SELL ONLINE IN GHANA",
        "What: Marketplace for businesses to list and sell products online.",
        "Who: Shops, suppliers, entrepreneurs with inventory to sell.",
        "How: Apply → Get approved → List products → Receive orders → Get paid (MoMo/bank).",
        "Offer now: " + ph("CURRENT PROMO SUMMARY"),
        "Apply: " + ph("SHOP APPLY URL") + "  |  WhatsApp: " + ph("YOUR WHATSAPP NUMBER"),
        "Contact: " + ph("YOUR FULL NAME") + " | " + ph("YOUR EMAIL"),
        "◆ Optional: small QR code to apply URL ◆",
    ]
    for i, line in enumerate(lines):
        one_pager.rows[i].cells[0].text = line
        if i == 0:
            set_cell_shading(one_pager.rows[i].cells[0], "E8F5E9")
    doc.add_paragraph()

    heading(doc, "Document C — Promoter / referral partner brief (1–2 pages)")
    add_insert_box(doc, "LOGO HERE — promoter brief header", "Optional header logo")
    bullet(doc, "What a promoter does: shares " + ph("REFERRAL / PROMOTER CODE") + ", introduces sellers, earns on paid registrations")
    bullet(doc, "Applicant benefit: " + ph("REFERRAL APPLICANT DISCOUNT"))
    bullet(doc, "Commission: " + ph("PROMOTER COMMISSION DESCRIPTION"))
    bullet(doc, "How to get a code: contact " + ph("YOUR EMAIL") + " or " + ph("YOUR PHONE"))
    bullet(doc, "Code of conduct: honest referrals, no spam, no misleading claims")

    heading(doc, "Document D — Meeting leave-behind (½ page)")
    body(doc, "Print or PDF after a call. Replace " + recipient("…") + " fields:")
    numbered(doc, "Thank you for your time today.")
    numbered(doc, "Recap: " + ph("COMPANY / BRAND NAME") + " helps " + recipient("Business name") + " sell online with " + recipient("Key benefit discussed") + ".")
    numbered(doc, "Next step: apply at " + ph("SHOP APPLY URL") + " or WhatsApp " + ph("YOUR WHATSAPP NUMBER") + " by " + recipient("Follow-up date") + ".")
    numbered(doc, "Referral code (if applicable): " + ph("REFERRAL / PROMOTER CODE"))
    numbered(doc, "Questions: " + ph("YOUR FULL NAME") + " — " + ph("YOUR PHONE"))

    add_word_layout_notes(doc)

    heading(doc, "DOCX formatting guide")
    add_table(
        doc,
        ["Element", "What you put there"],
        [
            ["Body font", "Calibri or Arial, 11pt"],
            ["Headings", "Word Heading 1 / 2; optional brand colour"],
            ["Margins", "2.5 cm all sides"],
            ["Header", insert_visual("SMALL LOGO") + " + document title"],
            ["Footer", "Page number + " + ph("YOUR WEBSITE URL")],
            ["Tables", "Light header row; edit cell text freely"],
            ["Appendix images", insert_visual("DASHBOARD SCREENSHOT") + ", " + insert_visual("SAMPLE PRODUCT LISTING")],
            ["Export", "File → Save As → PDF for WhatsApp"],
        ],
    )

    heading(doc, "Sample FAQ answers (replace placeholders)")
    faqs = [
        ("How long until my shop is approved?", "Typically " + ph("APPROVAL TIMEFRAME") + " after application and payment."),
        ("When do I get paid?", "Per platform payout rules — after order confirmation / fulfillment."),
        ("What products can I sell?", ph("PRODUCT CATEGORIES YOU ACCEPT")),
        ("Is there a registration fee?", ph("STANDARD REGISTRATION FEE") + " or FREE until " + ph("FREE REGISTRATION UNTIL DATE") + ". Code " + ph("REFERRAL / PROMOTER CODE") + " gives " + ph("REFERRAL APPLICANT DISCOUNT") + "."),
        ("Who do I contact for help?", ph("YOUR FULL NAME") + " — " + ph("YOUR EMAIL") + " / " + ph("YOUR PHONE")),
    ]
    add_table(doc, ["Question", "Sample answer"], faqs)

    path = OUT_DIR / "03_DPM_Business_Document_Samples.docx"
    doc.save(path)
    return path


def build_whatsapp_doc() -> Path:
    doc = Document()
    add_cover(
        doc,
        "WhatsApp & Social Messages",
        "Short copy-paste templates — personalise before sending",
    )

    heading(doc, "Rules for WhatsApp outreach")
    add_placeholder_legend(doc)
    bullet(doc, "Replace " + recipient("Recipient name") + " and " + recipient("Business name") + " every time")
    bullet(doc, "Attach one-pager PDF (from doc 03) after they reply — not on first cold ping unless they ask")

    messages = [
        ("Cold — short", "Hi " + recipient("Recipient name") + ", I’m " + ph("YOUR FULL NAME") + " from " + ph("COMPANY / BRAND NAME") + ". We help businesses like " + recipient("Business name") + " sell online. " + ph("CURRENT PROMO SUMMARY") + ". Interested? I can send details or " + ph("SHOP APPLY URL") + "."),
        ("Warm — short", "Hi " + recipient("Recipient name") + "! We’re onboarding sellers on " + ph("COMPANY / BRAND NAME") + " and I thought of you. List products, get orders, paid to MoMo/bank. " + ph("CURRENT PROMO SUMMARY") + ". Want the link?"),
        ("After no email reply", "Hi " + recipient("Recipient name") + ", sent you an email about " + ph("COMPANY / BRAND NAME") + " — just checking if online sales is something you’d consider this month. No pressure — happy to explain in 2 mins."),
        ("Send apply link", "Here’s your apply link: " + ph("APPLY URL WITH REF CODE") + ". Use code " + ph("REFERRAL / PROMOTER CODE") + " for " + ph("REFERRAL APPLICANT DISCOUNT") + ". Ping me on " + ph("YOUR WHATSAPP NUMBER") + " if the form gives trouble."),
        ("Promoter pitch", "Hi " + recipient("Recipient name") + ", we’re looking for referral partners for " + ph("COMPANY / BRAND NAME") + ". Share your code, new sellers get a discount, you earn " + ph("PROMOTER COMMISSION DESCRIPTION") + ". Interested?"),
        ("Post-meeting", "Thanks for the chat, " + recipient("Recipient name") + ". Next step: " + ph("SHOP APPLY URL") + ". I’m on WhatsApp if you need help listing your first product."),
    ]

    for title, text in messages:
        heading(doc, title, level=2)
        p = doc.add_paragraph()
        p.add_run(text).italic = True
        doc.add_paragraph()

    heading(doc, "Social post draft (optional)")
    add_insert_box(doc, "PROMO IMAGE OR BANNER (optional)", "Attach when posting on Facebook / Instagram / LinkedIn")
    body(
        doc,
        "🛍️ Now open for new sellers on "
        + ph("COMPANY / BRAND NAME")
        + "! List your products, reach customers online, get paid to MoMo or bank. "
        + ph("CURRENT PROMO SUMMARY")
        + ". Apply: "
        + ph("SHOP APPLY URL")
        + " #SellOnline #GhanaBusiness",
    )

    path = OUT_DIR / "04_DPM_WhatsApp_and_Social_Messages.docx"
    doc.save(path)
    return path


def build_checklist_doc() -> Path:
    doc = Document()
    add_cover(
        doc,
        "Outreach Checklist & Tracker",
        "Weekly tasks and simple metrics table",
    )

    heading(doc, "Before you start")
    body(doc, "Open " + ph("00_DPM_Where_To_Put_Your_Details.docx") + " first and fill in your master list.", bold=True)
    bullet(doc, "Add " + insert_visual("LOGO HERE") + " to covers and headers")
    bullet(doc, "Replace all " + ph("YOUR…") + " and fee placeholders in admin-aligned docs")
    bullet(doc, "Set " + ph("SHOP APPLY URL") + " and test " + ph("APPLY URL WITH REF CODE"))
    bullet(doc, "Export one-pager PDF from doc 03")
    bullet(doc, "Copy signature block into doc 02 emails")

    heading(doc, "Weekly checklist")
    for week, tasks in [
        ("Week 1", ["List 20 warm contacts", "Send 10 personalised messages", "Book 3 onboarding calls", "Goal: 3 applications"]),
        ("Week 2", ["Follow up non-replies", "Help pilots list first product", "Collect 1 testimonial", "Goal: 5 live shops"]),
        ("Week 3", ["Give referrers unique codes", "Send promoter template to 5 people", "Goal: 2 referral signups"]),
        ("Week 4", ["20 cold emails + 20 WhatsApp", "Review metrics table below", "Adjust offer if conversion is low"]),
    ]:
        heading(doc, week, level=2)
        for t in tasks:
            bullet(doc, t)

    heading(doc, "Simple tracker (copy to Excel or print)")
    add_table(
        doc,
        ["Date", "Contact", "Channel", "Template used", "Reply?", "Applied?", "Notes"],
        [["", "", "Email/WhatsApp/Call", "", "Y/N", "Y/N", ""] for _ in range(5)],
    )

    path = OUT_DIR / "05_DPM_Outreach_Checklist_and_Tracker.docx"
    doc.save(path)
    return path


def main() -> None:
    paths = [
        build_customization_guide_doc(),
        build_strategy_doc(),
        build_email_templates_doc(),
        build_business_samples_doc(),
        build_whatsapp_doc(),
        build_checklist_doc(),
    ]
    print("Generated:")
    for p in paths:
        print(f"  {p}")


if __name__ == "__main__":
    main()
