#!/usr/bin/env python3
"""Generate DanyPathMart Hostinger step-by-step deployment guide (Word)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "DanyPathMart-Hostinger-Deployment-Guide.docx"
DOMAIN = "danypathmart.store"


def add_title(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph("")


def build() -> Document:
    doc = Document()
    today = date.today().strftime("%B %d, %Y")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("DanyPathMart\n")
    r.bold = True
    r.font.size = Pt(26)
    r2 = title.add_run("Hostinger Deployment Guide\n")
    r2.font.size = Pt(15)
    r3 = title.add_run(f"Step-by-step · {today}")
    r3.font.size = Pt(11)
    doc.add_paragraph("")
    add_para(
        doc,
        f"This guide walks you through hosting DanyPathMart on Hostinger at https://{DOMAIN}. "
        "Stack: React (Vite) storefront in public_html, PHP 8.2 API in public_html/api, MySQL database.",
    )
    doc.add_page_break()

    add_title(doc, "Contents", 1)
    add_bullets(
        doc,
        [
            "1. Before you start",
            "2. Domain & Hostinger setup",
            "3. Create the MySQL database",
            "4. Build the frontend on your computer",
            "5. Prepare the backend (Composer)",
            "6. Upload files to Hostinger",
            "7. Configure production environment (.env)",
            "8. Run database migrations",
            "9. PHP, SSL, and upload folders",
            "10. Paystack & email",
            "11. Go-live checks",
            "12. After launch (admin setup)",
            "13. Troubleshooting",
        ],
    )
    doc.add_page_break()

    # 1
    add_title(doc, "1. Before you start", 1)
    add_para(doc, "You will need:", bold=True)
    add_bullets(
        doc,
        [
            "Hostinger account with your domain pointed to Hostinger",
            "Node.js and npm on your PC (to build the React app)",
            "PHP 8.2+ and Composer on your PC (or run Composer on the server via SSH)",
            "Paystack live API keys (for real payments)",
            "Hostinger email account (e.g. noreply@danypathmart.store) for SMTP",
            "Project files from the DanyPathMart repository",
        ],
    )
    add_para(doc, "Production URLs:", bold=True)
    add_bullets(
        doc,
        [
            f"Storefront: https://{DOMAIN}",
            f"API: https://{DOMAIN}/api",
            f"Admin: https://{DOMAIN}/admin",
            f"Health check: https://{DOMAIN}/api/health",
        ],
    )

    # 2
    add_title(doc, "2. Domain & Hostinger setup", 1)
    add_numbered(
        doc,
        [
            "Log in to Hostinger hPanel.",
            f"Add or connect the website for {DOMAIN} (Websites → Add website or manage existing).",
            "Point your domain nameservers to Hostinger (at your domain registrar) if not already done.",
            "Wait for DNS to propagate (can take up to 24–48 hours; often faster).",
            "Open File Manager and note the folder public_html — this is your site root.",
        ],
    )

    # 3
    add_title(doc, "3. Create the MySQL database", 1)
    add_numbered(
        doc,
        [
            "hPanel → Databases → MySQL Databases.",
            "Create a new database named danypathmart.",
            "Create a MySQL user with a strong password.",
            "Assign the user to the database with All Privileges.",
            "Write down: DB host (usually localhost), full database name, username, password.",
        ],
    )
    add_para(
        doc,
        "Important: On shared hosting, Hostinger may prefix names (e.g. u123456789_danypathmart). "
        "Use the full prefixed name in DB_NAME in your .env file.",
    )
    add_para(doc, "Import base schema (phpMyAdmin):", bold=True)
    add_numbered(
        doc,
        [
            "hPanel → phpMyAdmin → select your database.",
            "Import database/schema.sql from the project.",
            "Import database/seed.sql (creates super admin — change password before go-live).",
        ],
    )
    add_para(doc, "Default seed admin (change immediately after first login):", bold=True)
    add_bullets(
        doc,
        [
            "Email: admin@danypathmart.store",
            "Password: see database/seed.sql comments (default in seed file)",
        ],
    )

    # 4
    add_title(doc, "4. Build the frontend on your computer", 1)
    add_numbered(
        doc,
        [
            "Open a terminal in the project folder.",
            "cd frontend",
            "Copy .env.production.example to .env.production",
            "Edit .env.production with your live values (see Section 7.2).",
            "Run: npm ci",
            "Run: npm run build",
            "When finished, built files are in frontend/dist/",
        ],
    )
    add_para(doc, "Do NOT set VITE_DEV_ADMIN_BYPASS in production builds.")

    # 5
    add_title(doc, "5. Prepare the backend (Composer)", 1)
    add_numbered(
        doc,
        [
            "cd backend (from project root)",
            "Run: composer install --no-dev --optimize-autoloader",
            "This creates the vendor/ folder required on the server.",
        ],
    )
    add_para(
        doc,
        "If you cannot run Composer on your PC, upload backend without vendor/ and run "
        "composer install --no-dev on the server via SSH (see Section 6).",
    )

    # 6
    add_title(doc, "6. Upload files to Hostinger", 1)
    add_para(doc, "Use File Manager or FTP (e.g. FileZilla). Upload mapping:", bold=True)
    add_table(
        doc,
        ["From your PC", "To Hostinger"],
        [
            ["frontend/dist/* (all files inside dist)", "public_html/"],
            ["deploy/public_html.htaccess", "public_html/.htaccess"],
            ["Entire backend/ folder", "public_html/api/"],
            ["frontend/public/robots.txt", "public_html/robots.txt (if not in dist)"],
        ],
    )
    add_para(doc, "Final folder layout on the server:", bold=True)
    add_code(
        doc,
        "public_html/\n"
        "├── .htaccess\n"
        "├── index.html\n"
        "├── assets/\n"
        "├── robots.txt\n"
        "└── api/\n"
        "    ├── .htaccess\n"
        "    ├── .env          (you create in next section)\n"
        "    ├── index.php\n"
        "    ├── vendor/\n"
        "    └── uploads/      (create in Section 9)",
    )

    # 7
    add_title(doc, "7. Configure production environment (.env)", 1)
    add_title(doc, "7.1 Backend — public_html/api/.env", 2)
    add_para(doc, "Copy backend/.env.production.example to public_html/api/.env and fill in:")
    add_code(
        doc,
        f"APP_ENV=production\n"
        f"APP_URL=https://{DOMAIN}/api\n"
        f"CORS_ORIGIN=https://{DOMAIN}\n"
        f"ADMIN_EMAIL=admin@{DOMAIN}\n"
        f"DEV_ADMIN_BYPASS=0\n\n"
        f"DB_HOST=localhost\n"
        f"DB_NAME=danypathmart\n"
        f"DB_USER=your_db_user\n"
        f"DB_PASS=your_db_password\n\n"
        f"JWT_SECRET=<64-char random hex>\n"
        f"PAYSTACK_SECRET_KEY=sk_live_...\n"
        f"PAYSTACK_PUBLIC_KEY=pk_live_...\n\n"
        f"SMTP_HOST=smtp.hostinger.com\n"
        f"SMTP_PORT=587\n"
        f"SMTP_USER=noreply@{DOMAIN}\n"
        f"SMTP_PASS=...\n"
        f"SMTP_FROM_NAME=DanyPathMart\n\n"
        f"WEBAUTHN_RP_ID={DOMAIN}\n"
        f"WEBAUTHN_RP_NAME=DanyPathMart\n"
        f"WEBAUTHN_ORIGIN=https://{DOMAIN}",
    )
    add_para(doc, "Generate JWT_SECRET on your PC:", bold=True)
    add_code(doc, 'php -r "echo bin2hex(random_bytes(32));"')
    add_para(
        doc,
        "WEBAUTHN_RP_ID must match your domain exactly (no www, no https://, no path).",
    )

    add_title(doc, "7.2 Frontend — frontend/.env.production (before build)", 2)
    add_code(
        doc,
        f"VITE_API_BASE_URL=https://{DOMAIN}/api\n"
        f"VITE_APP_NAME=DanyPathMart\n"
        f"VITE_PAYSTACK_PUBLIC_KEY=pk_live_...\n"
        f"VITE_WEBAUTHN_RP_ID={DOMAIN}",
    )
    add_para(doc, "Rebuild (npm run build) after changing .env.production, then re-upload public_html files.")

    add_title(doc, "7.3 Verify environment", 2)
    add_numbered(
        doc,
        [
            "If you have SSH: cd public_html/api && php scripts/check-production-env.php",
            f"Open https://{DOMAIN}/api/.env in a browser — must return 403 Forbidden (not file contents).",
        ],
    )

    # 8
    add_title(doc, "8. Run database migrations", 1)
    add_para(
        doc,
        "Schema + seed are not enough. The app needs all migration files (001 through 048+).",
    )
    add_numbered(
        doc,
        [
            "Ensure public_html/api/.env points to your Hostinger database.",
            "SSH into Hostinger (or run locally with .env pointing at remote DB).",
            "cd public_html/api",
            "Run: php scripts/migrate-all.php",
            "Confirm no fatal errors; duplicate column messages are OK on re-run.",
        ],
    )
    add_para(doc, "Take a full database backup in hPanel before running migrations on an existing live site.")

    # 9
    add_title(doc, "9. PHP, SSL, and upload folders", 1)
    add_numbered(
        doc,
        [
            "hPanel → PHP Configuration → select PHP 8.2 or newer for your domain.",
            f"hPanel → SSL → install Let's Encrypt certificate for {DOMAIN}.",
            "Create upload directories (SSH or File Manager):",
        ],
    )
    add_code(
        doc,
        "cd public_html/api\n"
        "mkdir -p uploads/products uploads/searches uploads/avatars\n"
        "mkdir -p uploads/careers uploads/shops uploads/customizations uploads/hero\n"
        "chmod 755 uploads uploads/*",
    )
    add_para(doc, "These folders must be writable so product images and uploads work.")

    # 10
    add_title(doc, "10. Paystack & email", 1)
    add_title(doc, "10.1 Paystack", 2)
    add_numbered(
        doc,
        [
            "Log in to Paystack Dashboard (live mode).",
            f"Settings → API Keys & Webhooks → Webhook URL: https://{DOMAIN}/api/payments/webhook",
            "Copy live secret key → PAYSTACK_SECRET_KEY in api/.env",
            "Copy live public key → PAYSTACK_PUBLIC_KEY in api/.env and VITE_PAYSTACK_PUBLIC_KEY in frontend build",
        ],
    )
    add_title(doc, "10.2 Email (Hostinger SMTP)", 2)
    add_numbered(
        doc,
        [
            f"hPanel → Emails → create noreply@{DOMAIN} (or use existing mailbox).",
            "Use smtp.hostinger.com, port 587, in api/.env.",
            "Test: trigger password reset or registration email from the live site.",
        ],
    )

    # 11
    add_title(doc, "11. Go-live checks", 1)
    add_bullets(
        doc,
        [
            f"https://{DOMAIN} loads the storefront",
            f"https://{DOMAIN}/api/health returns success with database connected",
            f"https://{DOMAIN}/admin — log in as super admin",
            "Change default admin password immediately",
            "Add at least one product in Admin → Products",
            "Test checkout with a small Paystack payment (live or test mode as appropriate)",
            "Confirm Paystack webhook shows successful delivery in dashboard",
        ],
    )
    add_para(doc, "Optional smoke tests (from your PC with project repo):", bold=True)
    add_code(
        doc,
        f"php backend/scripts/pilot-smoke-test.php https://{DOMAIN}/api\n"
        f"php backend/scripts/p5-smoke-test.php https://{DOMAIN}/api",
    )

    # 12
    add_title(doc, "12. After launch (admin setup)", 1)
    add_numbered(
        doc,
        [
            "Admin → Company settings: business name, phone, WhatsApp, return policy, payment methods.",
            "Admin → Hero banners: add homepage slides (no demo banners on fresh install).",
            "Admin → Launch readiness: review P1 checks; apply pilot preset if starting small.",
            "Enable optional modules only when ready: marketplace, pickup stations, drivers, etc.",
            "Enable 2FA on super admin account (recommended).",
        ],
    )

    # 13
    add_title(doc, "13. Troubleshooting", 1)
    add_table(
        doc,
        ["Problem", "What to check"],
        [
            ["Blank page or 404 on /shop", ".htaccess in public_html; index.html uploaded"],
            ["API errors / CORS", "CORS_ORIGIN=https://danypathmart.store (no trailing slash); APP_URL ends with /api"],
            ["Database connection failed", "DB_NAME uses Hostinger prefixed name; user has privileges"],
            ["Payments not confirming", "Paystack webhook URL; PAYSTACK_SECRET_KEY is live sk_ key"],
            ["Images won't upload", "uploads/ folders exist and are writable (chmod 755)"],
            [".env visible in browser", "api/.htaccess must block .env; contact Hostinger if 200 returned"],
            ["White screen after deploy", "PHP 8.2+; check Hostinger error logs in hPanel"],
        ],
    )
    add_para(doc, "For full security checklist and rollback steps, see DEPLOY.md and LAUNCH.md in the repository.")

    doc.add_paragraph("")
    add_para(doc, f"Document generated for DanyPathMart · https://{DOMAIN}", bold=False)
    return doc


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
