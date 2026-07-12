#!/usr/bin/env python3
"""Generate DanyPathMart master project documentation (Word)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "DanyPathMart-Project-Documentation.docx"


def add_title(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph("")


def build() -> Document:
    doc = Document()
    today = date.today().strftime("%B %d, %Y")

    # Cover
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("DanyPathMart\n")
    r.bold = True
    r.font.size = Pt(28)
    r2 = title.add_run("Complete Project Documentation\n")
    r2.font.size = Pt(16)
    r3 = title.add_run(f"Version 1.0 · {today}")
    r3.font.size = Pt(11)
    doc.add_paragraph("")
    add_para(
        doc,
        "Supermarket and marketplace — e-commerce platform "
        "for Ghana (GHS). Stack: React (Vite) storefront + PHP 8.2 API + MySQL.",
    )
    doc.add_page_break()

    # TOC placeholder note
    add_title(doc, "Document contents", 1)
    add_bullets(
        doc,
        [
            "1. Project overview",
            "2. Roles, duties & permissions",
            "3. System architecture & folder layout",
            "4. Local development setup",
            "5. Database installation & migrations",
            "6. Hostinger hosting & deployment",
            "7. Environment variables",
            "8. Launch phases (P0–P5)",
            "9. Admin modules reference",
            "10. CLI scripts & smoke tests",
            "11. Security checklist",
            "12. Ongoing operations",
            "13. Planned features",
            "14. Troubleshooting & rollback",
        ],
    )
    doc.add_page_break()

    # 1 Overview
    add_title(doc, "1. Project overview", 1)
    add_para(doc, "DanyPathMart is a full-stack online store with:")
    add_bullets(
        doc,
        [
            "Public storefront: shop, cart, checkout (Paystack MoMo/card), customer accounts",
            "Admin Seller Center: products, orders, staff, company settings, launch readiness",
            "Optional modules: pickup stations, hub/driver logistics, station repack, marketplace sellers",
            "Workforce HR: employee profiles, leave requests",
            "Storefront CMS: hero banners, legal policies",
            "Careers & driver hiring portals",
        ],
    )
    add_para(doc, "Primary URLs (production example):", bold=True)
    add_bullets(
        doc,
        [
            "Storefront: https://danypathmart.store",
            "API: https://danypathmart.store/api",
            "Admin: https://danypathmart.store/admin",
            "Health check: https://danypathmart.store/api/health",
        ],
    )

    # 2 Roles
    add_title(doc, "2. Roles, duties & permissions", 1)
    add_title(doc, "2.1 User roles", 2)
    add_table(
        doc,
        ["Role", "Portal", "Main duties"],
        [
            ["Super admin", "/admin", "Full access; company settings; module toggles; staff & permissions"],
            ["Staff", "/admin", "Day-to-day ops per assigned permissions (orders, products, etc.)"],
            ["Customer", "/dashboard", "Orders, wallet, wishlist, addresses, quotes"],
            ["Driver", "/driver", "Delivery runs assigned from hub logistics"],
            ["Station staff", "/station", "Repack orders for pickup stations"],
            ["Shop seller", "/seller", "Marketplace shop dashboard (when approved)"],
        ],
    )

    add_title(doc, "2.2 Recommended staff duties", 2)
    add_table(
        doc,
        ["Function", "Typical permissions", "Daily duties"],
        [
            ["Store manager", "view/edit orders, products, shipping", "Fulfill orders, update statuses, stock"],
            ["Catalogue editor", "view/add_edit products, categories, kits", "SKUs, prices, images, categories"],
            ["Customer support", "view_users, contact inbox, view_orders", "Inbox, order lookups, WhatsApp follow-up"],
            ["Finance", "view_reports, export_reports, view_shop_billing (planned)", "Reconciliation, Paystack, payouts"],
            ["HR", "view/manage employees, leave requests", "Profiles, leave approve/reject"],
            ["Logistics lead", "pickup, hub, delivery runs, station staff", "Stations, drivers, runs"],
            ["Marketplace ops", "manage_marketplace, approve_shop_listings", "Applications, listings, withdrawals"],
            ["Marketing / content", "hero banners, legal policies", "Homepage, policies, footer trust"],
        ],
    )

    add_title(doc, "2.3 Permission groups (RBAC)", 2)
    add_para(doc, "Assign under Admin → Position access. Super admin has all permissions implicitly.")
    groups = [
        ("Orders", "view_orders, edit_orders"),
        ("Quotes", "view_quotes, edit_quotes"),
        ("Products & catalogue", "view/add/edit/delete products, categories, size guides, kits, flash sale"),
        ("Customers & comms", "view/edit users, notifications, contact inbox"),
        ("Staff & hiring", "manage_staff, hire_employees, position permissions, employees, leave"),
        ("Careers", "manage_careers"),
        ("Marketplace", "manage_marketplace, approve_shop_listings, view_shop_billing, manage_shop_fees, waive_shop_fees"),
        ("Storefront content", "hero banners, legal policies"),
        ("Company & shipping", "company settings, shipping, pickup, hub, drivers, station staff"),
        ("Reports", "view_reports, export_reports"),
        ("Image alerts", "view/manage image alerts"),
        ("Custom proofs", "manage_custom_proofs"),
    ]
    add_table(doc, ["Group", "Permission keys (summary)"], [[a, b] for a, b in groups])

    # 3 Architecture
    add_title(doc, "3. System architecture & folder layout", 1)
    add_title(doc, "3.1 Repository structure", 2)
    layout = """
danypathmart/
├── frontend/          React + Vite SPA (storefront + admin UI)
├── backend/           PHP API (index.php router, api/, helpers/)
├── database/          schema.sql, seed.sql, migrations/
├── deploy/            Hostinger .htaccess templates
├── docs/              Extended documentation
├── LAUNCH.md          Launch checklist P0–P5
├── DEPLOY.md          Deploy & security
└── README.md          Quick start
"""
    doc.add_paragraph(layout)

    add_title(doc, "3.2 Production layout on Hostinger", 2)
    host_layout = """
public_html/
├── .htaccess              SPA rewrite
├── index.html, assets/    from frontend/dist/
├── robots.txt, sitemap.xml
└── api/
    ├── .htaccess
    ├── .env               secrets (must return 403 in browser)
    ├── index.php
    ├── vendor/            Composer dependencies
    └── uploads/           products, avatars, careers, shops, hero, etc.
"""
    doc.add_paragraph(host_layout)

    # 4 Local dev
    add_title(doc, "4. Local development setup", 1)
    add_title(doc, "4.1 Prerequisites", 2)
    add_bullets(
        doc,
        [
            "PHP 8.2+ with extensions: pdo_mysql, mbstring, openssl, json",
            "Composer",
            "Node.js 18+ and npm",
            "MySQL 8 / MariaDB",
        ],
    )
    add_title(doc, "4.2 Backend", 2)
    add_numbered(
        doc,
        [
            "Copy backend/.env.example → backend/.env and fill DB + JWT",
            "cd backend && composer install",
            "Create MySQL database; import database/schema.sql and database/seed.sql",
            "php backend/scripts/migrate-all.php",
            "php -S localhost:8000 -t backend (or point Apache/Nginx to backend/)",
        ],
    )
    add_title(doc, "4.3 Frontend", 2)
    add_numbered(
        doc,
        [
            "Copy frontend/.env.example → frontend/.env",
            "Set VITE_API_BASE_URL=http://localhost:8000",
            "cd frontend && npm ci && npm run dev",
            "Open http://localhost:5173",
        ],
    )
    add_para(doc, "Default seed super admin: see database/seed.sql — change password immediately.")

    # 5 Database
    add_title(doc, "5. Database installation & migrations", 1)
    add_title(doc, "5.1 Fresh database", 2)
    add_numbered(
        doc,
        [
            "Create empty MySQL database",
            "Import database/schema.sql",
            "Import database/seed.sql (optional on production if admin created manually)",
            "Run: php backend/scripts/migrate-all.php",
        ],
    )
    add_para(doc, "Never deploy with schema + seed only — migrations through 046 are required.")
    add_title(doc, "5.2 Key migration phases", 2)
    add_table(
        doc,
        ["Range", "Summary"],
        [
            ["001–013", "Core commerce: checkout, sizing, kits, group orders"],
            ["030–033", "Careers & hiring"],
            ["034–040", "Pickup, marketplace, referrals, hub logistics, station repack"],
            ["041–042", "Employee staff IDs (DPM-EMP-####)"],
            ["043–044", "Hero banners & legal policies CMS"],
            ["045", "HR: leave requests, employee profiles"],
            ["046", "Ops: analytics toggle, uptime monitor URL"],
        ],
    )
    add_para(doc, "Before production migrate-all: take a full MySQL backup (hPanel or mysqldump).")

    # 6 Hostinger
    add_title(doc, "6. Hostinger hosting & deployment", 1)
    add_title(doc, "6.1 Pre-deploy checklist", 2)
    add_bullets(
        doc,
        [
            "Domain pointed to Hostinger; SSL (Let's Encrypt) enabled",
            "MySQL database created in hPanel as danypathmart (DB_NAME=danypathmart; use Hostinger prefixed name if shown)",
            "Paystack live keys and webhook URL ready",
            "SMTP credentials for transactional email",
        ],
    )

    add_title(doc, "6.2 Build frontend (on your PC)", 2)
    add_numbered(
        doc,
        [
            "cd frontend",
            "Copy .env.production.example → .env.production",
            "Set VITE_API_BASE_URL=https://danypathmart.store/api",
            "Set VITE_PAYSTACK_PUBLIC_KEY=pk_live_...",
            "Set VITE_WEBAUTHN_RP_ID=yourdomain.com",
            "Do NOT set VITE_DEV_ADMIN_BYPASS",
            "npm ci && npm run build",
        ],
    )

    add_title(doc, "6.3 Upload to Hostinger (File Manager or FTP)", 2)
    add_table(
        doc,
        ["Local path", "Upload to"],
        [
            ["frontend/dist/*", "public_html/ (root)"],
            ["deploy/public_html.htaccess", "public_html/.htaccess"],
            ["backend/* (entire folder)", "public_html/api/"],
        ],
    )
    add_para(doc, "Include backend/vendor/ OR run on server: cd public_html/api && composer install --no-dev")

    add_title(doc, "6.4 Server configuration", 2)
    add_numbered(
        doc,
        [
            "hPanel → PHP → select PHP 8.2 or newer for the domain",
            "Create public_html/api/.env from backend/.env.production.example",
            "Run: php scripts/check-production-env.php (via SSH or local copy of .env)",
            "Verify https://danypathmart.store/api/.env returns 403 Forbidden",
            "Create writable upload folders (see DEPLOY.md): uploads/products, searches, avatars, careers, shops, customizations, hero — chmod 755",
            "Paystack dashboard → Webhook → https://danypathmart.store/api/payments/webhook",
        ],
    )

    add_title(doc, "6.5 Database on production", 2)
    add_numbered(
        doc,
        [
            "hPanel → phpMyAdmin → Import schema.sql then seed.sql (or restore backup)",
            "SSH: cd public_html/api && php scripts/migrate-all.php",
            "Log in to admin; change super admin password",
            "Admin → Launch readiness → run checks",
        ],
    )

    add_title(doc, "6.6 Post-upload smoke tests", 2)
    add_bullets(
        doc,
        [
            "https://danypathmart.store loads storefront",
            "https://danypathmart.store/api/health returns db:true",
            "Admin login works",
            "php backend/scripts/pilot-smoke-test.php https://danypathmart.store/api",
            "php backend/scripts/p5-smoke-test.php https://danypathmart.store/api",
        ],
    )

    # 7 Env
    add_title(doc, "7. Environment variables", 1)
    add_title(doc, "7.1 Backend (public_html/api/.env)", 2)
    add_table(
        doc,
        ["Variable", "Purpose"],
        [
            ["APP_ENV", "production on live server"],
            ["APP_URL", "https://domain.com/api"],
            ["CORS_ORIGIN", "https://domain.com (no trailing slash)"],
            ["JWT_SECRET", "Long random hex — never reuse dev value"],
            ["DEV_ADMIN_BYPASS", "Must be 0 in production"],
            ["DB_NAME", "danypathmart (or u123456789_danypathmart on Hostinger shared hosting)"],
            ["DB_USER / DB_PASS", "MySQL user with full privileges on that database"],
            ["PAYSTACK_SECRET_KEY", "sk_live_... for production"],
            ["SMTP_*", "Email for password reset, notifications"],
            ["WEBAUTHN_RP_ID", "Apex domain e.g. danypathmart.store"],
            ["GOOGLE_VISION_API_KEY", "Image search (server only)"],
            ["SENTRY_DSN", "Optional error monitoring (P5)"],
        ],
    )
    add_title(doc, "7.2 Frontend build (.env.production)", 2)
    add_table(
        doc,
        ["Variable", "Purpose"],
        [
            ["VITE_API_BASE_URL", "https://domain.com/api"],
            ["VITE_PAYSTACK_PUBLIC_KEY", "pk_live_..."],
            ["VITE_WEBAUTHN_RP_ID", "Same as backend WEBAUTHN_RP_ID"],
        ],
    )

    # 8 Launch phases
    add_title(doc, "8. Launch phases (P0–P5)", 1)
    add_title(doc, "8.1 Phase P0 — Foundation", 2)
    add_para(doc, "Blocking before any public URL: DB migrated, production .env, SSL, uploads writable, admin password changed.")
    add_title(doc, "8.2 Business Phase 1 — Go live safely", 2)
    add_bullets(
        doc,
        [
            "Set PAYSTACK_SECRET_KEY; complete live test payment + webhook",
            "Apply pilot preset (turns off unneeded modules): Admin → Launch readiness or php backend/scripts/pilot-preset.php",
            "Complete P1 + P2 automated checks in /admin/launch-readiness",
            "3+ real pilot orders fulfilled",
        ],
    )
    add_title(doc, "8.3 Business Phase 2 — Operate and grow", 2)
    add_bullets(
        doc,
        [
            "Enable modules one at a time from Launch readiness → Phase 2",
            "HR (P4) and analytics (P5) when staffed",
            "Weekly ops checklist: webhooks, logs, smoke tests",
        ],
    )
    add_title(doc, "8.4 Technical phases P1–P5", 2)
    add_table(
        doc,
        ["Phase", "Focus"],
        [
            ["P1", "Pilot commerce: catalog, Paystack, company info, shipping"],
            ["P2", "Public launch: legal policies, footer, checkout trust, SEO"],
            ["P3", "Logistics & marketplace (conditional on toggles)"],
            ["P4", "Workforce HR: employees, leave"],
            ["P5", "Quality & ops: smoke tests, GA4, Sentry, uptime"],
        ],
    )

    # 9 Admin modules
    add_title(doc, "9. Admin modules reference", 1)
    add_table(
        doc,
        ["Area", "Path", "Notes"],
        [
            ["Dashboard", "/admin/dashboard", "Stats, launch banners"],
            ["Launch readiness", "/admin/launch-readiness", "P0–P5 checks, pilot preset, module rollout"],
            ["Orders & alerts", "/admin/orders, /admin/alerts", "Order management"],
            ["Products", "/admin/products, categories, kits, size guides", "Catalogue"],
            ["Customers", "/admin/users, contact-inbox", "Accounts & messages"],
            ["Shipping", "/admin/shipping", "Delivery % and rules"],
            ["Pickup / Hub / Runs", "/admin/pickup-stations, hub-logistics, delivery-runs", "P3 logistics"],
            ["Station staff", "/admin/station-staff", "Repack portal accounts"],
            ["Marketplace", "/admin/marketplace", "Applications, shops, listings, withdrawals"],
            ["Hero & legal", "/admin/hero-banners, /admin/legal-policies", "Storefront CMS"],
            ["Employees & leave", "/admin/employees, /admin/leave-requests", "P4 HR"],
            ["Staff & permissions", "/admin/staff, /admin/position-permissions", "RBAC"],
            ["Company settings", "/admin/company-settings", "Contact, modules, Paystack toggle, HR, analytics"],
            ["Reports", "/admin/reports", "Sales and exports"],
        ],
    )

    # 10 CLI
    add_title(doc, "10. CLI scripts & smoke tests", 1)
    add_table(
        doc,
        ["Command", "Purpose"],
        [
            ["php backend/scripts/migrate-all.php", "Apply all DB migrations"],
            ["php backend/scripts/check-production-env.php", "Pre-flight .env validation"],
            ["php backend/scripts/launch-readiness-cli.php", "All P1–P5 automated checks"],
            ["php backend/scripts/pilot-smoke-test.php [API_URL]", "Basic API health smoke test"],
            ["php backend/scripts/p5-smoke-test.php [API_URL]", "Extended API smoke test"],
            ["php backend/scripts/pilot-preset.php", "Disable extended modules for pilot"],
            ["php -r \"echo bin2hex(random_bytes(32));\"", "Generate JWT_SECRET"],
        ],
    )

    # 11 Security
    add_title(doc, "11. Security checklist", 1)
    add_bullets(
        doc,
        [
            ".env blocked by .htaccess (403 in browser)",
            "DEV_ADMIN_BYPASS off in production",
            "JWT_SECRET unique and long; super admin password changed from seed",
            "Paystack webhook uses hash_equals for signature verification",
            "PDO prepared statements for user input",
            "Staff cannot edit own permissions",
            "Super admin account cannot be deleted via API",
            "GOOGLE_VISION_API_KEY never exposed to frontend",
            "2FA recommended for super admin before marketing",
            "WEBAUTHN_RP_ID matches production domain exactly",
        ],
    )

    # 12 Ops
    add_title(doc, "12. Ongoing operations", 1)
    add_title(doc, "12.1 Daily / weekly", 2)
    add_bullets(
        doc,
        [
            "Monitor admin order alerts and contact inbox",
            "Review Paystack dashboard for failed webhooks",
            "Check PHP error logs on Hostinger after deploys",
            "Run p5-smoke-test after each production deploy",
        ],
    )
    add_title(doc, "12.2 Backups", 2)
    add_bullets(
        doc,
        [
            "Schedule daily MySQL backup in hPanel",
            "Backup uploads/ folder periodically",
            "Before migrate-all on production: manual mysqldump",
        ],
    )
    add_title(doc, "12.3 Module toggles", 2)
    add_para(doc, "Admin → Company settings → Platform modules. Enable only what you operate:")
    add_table(
        doc,
        ["Module", "Enable when"],
        [
            ["Pickup stations", "Stations exist and staff trained"],
            ["Driver / hub logistics", "Hub and drivers ready"],
            ["Station repack", "Station portal in use"],
            ["Marketplace", "Sellers ready to onboard"],
            ["Careers", "Hiring actively"],
            ["Leave requests (HR)", "HR process defined"],
        ],
    )

    # 13 Planned
    add_title(doc, "13. Planned features", 1)
    add_para(doc, "Marketplace shop billing (documented, not built): registration fee on apply + monthly/yearly renewal set by admin. Permissions: view_shop_billing, manage_shop_fees, waive_shop_fees. See docs/MARKETPLACE_SHOP_BILLING.md in repository.")

    # 14 Troubleshooting
    add_title(doc, "14. Troubleshooting & rollback", 1)
    add_table(
        doc,
        ["Issue", "Action"],
        [
            ["API 500", "Check api/.env, PHP version, error logs; verify vendor/ installed"],
            ["CORS errors", "CORS_ORIGIN must match storefront URL exactly"],
            ["Paystack webhook not updating orders", "Verify webhook URL and live secret key"],
            ["Admin blank after deploy", "Rebuild frontend with correct VITE_API_BASE_URL"],
            ["Uploads fail", "Check uploads/ folder permissions (755)"],
            ["Migration errors", "Restore DB backup; fix SQL; re-run migrate-all"],
        ],
    )
    add_title(doc, "14.1 Rollback procedure", 2)
    add_numbered(
        doc,
        [
            "Restore MySQL backup from before the failed change",
            "Re-upload previous frontend/dist and api/ snapshot",
            "Keep Paystack webhook URL aligned with API base URL",
        ],
    )

    # Footer
    doc.add_page_break()
    add_para(doc, "— End of document —", bold=True)
    add_para(doc, "Maintained with the DanyPathMart codebase. Update this file by running:")
    add_para(doc, "python backend/scripts/generate-project-documentation.py")

    return doc


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build()
    doc.save(str(OUT))
    print(f"Written: {OUT}")


if __name__ == "__main__":
    main()
